"""เก็บโพสต์ของ **เจ้าของบัญชีเอง** ทั้งหมดจากกลุ่มที่ระบุ แล้วแยกแมส/ไม่แมส

สั่งโดยเจ้าของโปรเจกต์ 23 ส.ค. 2569 (แก้จากรอบแรก):
  "ให้เก็บโพสต์ทั้งหมดเลย แต่แยกโพสต์ที่แมส (เกิน 100) และไม่แมส อยู่คนละ
   group กันเพื่อเอามาวิเคราะห์ · และให้แยกคอนเทนต์ตามสินค้าด้วย"

**ทำไมการเก็บทุกโพสต์ถึงถูกกว่าเก็บแต่โพสต์ดัง**
ดูแต่โพสต์ที่ดัง แล้วสรุปว่า "โพสต์ดังมีราคาอยู่ในแคปชัน 80%" — ฟังดูเป็นสูตร
แต่ถ้าโพสต์ที่ไม่ดังก็มีราคา 80% เหมือนกัน แปลว่าราคา **ไม่ใช่** ตัวแปรเลย
ต้องมีกลุ่มเทียบเท่านั้นถึงจะแยกออกว่าอะไรคือสูตรจริง อะไรคือเรื่องบังเอิญ

**สินค้าจับกลุ่มจากตัวข้อความ** โพสต์แมสคือเอาเนื้อเดียวกันไปลงหลายกลุ่ม
โพสต์ที่แคปชันเหมือนกันจึงเป็นสินค้าเดียวกัน — และได้ของแถมสำคัญ:
เนื้อเดียวกันลงคนละกลุ่ม **กลุ่มไหนตอบดีกว่า** ซึ่งเป็นคำถามหลักของโพสต์แมส

ใช้ของเดิมในโปรเจกต์เกือบทั้งหมด
  fb_posts_parse.FeedParser      แกะ Story จาก GraphQL
  fb_posts_comments.collect_comments  เก็บคอมเมนต์ครบทุกอันรวมคำตอบซ้อน
  fb_posts_collect.fetch_image   โหลดรูปจาก CDN
  fb_mass_finder.launch_bot_browser   เปิด Chrome ของบอทพร้อมถือล็อก

รัน:
    python fb_myposts.py                       เก็บครบ + คอมเมนต์
    python fb_myposts.py --no-comments         เก็บแค่โพสต์ (เร็วมาก)
    python fb_myposts.py --mass 100            เปลี่ยนเส้นแบ่งแมส
    python fb_myposts.py --comments-min 1      เปิดโพสต์ที่มีคอมเมนต์ตั้งแต่กี่อัน
    python fb_myposts.py --dry-run             ตรวจของก่อนเปิดเบราว์เซอร์
"""
from __future__ import annotations

import argparse
import json
import re
import sys
import time
import traceback
from datetime import datetime
from pathlib import Path

import fb_mass_finder as mf
import fb_myposts_features as feat
import fb_posts_collect as pc
import fb_posts_comments as pcm
import fb_posts_parse as pp
from bot_profiles import ProfileFarm

HERE = Path(__file__).resolve().parent
OUT_DIR = mf.DATA_DIR / "myposts"
IMAGE_DIR = OUT_DIR / "images"
COMMENT_IMAGE_DIR = OUT_DIR / "comment_images"
RAW_FILE = OUT_DIR / "myposts.json"
XLSX_FILE = OUT_DIR / "myposts.xlsx"
GROUPS_FILE = mf.DATA_DIR / "myposts_groups.txt"
PRODUCTS_FILE = mf.DATA_DIR / "myposts_products.txt"
LOG_FILE = mf.DATA_DIR / "logs" / "fb_myposts.log"

DEFAULT_BOT = "Bot7"
MASS_MIN = 100            # ไลค์+คอมเมนต์+แชร์ "เกิน" ค่านี้ = แมส
COMMENTS_MIN = 1          # เปิดโพสต์เก็บคอมเมนต์เมื่อมีอย่างน้อยกี่อัน

DEFAULT_GROUPS = [
    "https://www.facebook.com/share/g/1Eap6QKBGy/",
    "https://www.facebook.com/share/g/1Ez3SvERNU/",
    "https://www.facebook.com/share/g/14vh4i7WU13/",
    "https://www.facebook.com/share/g/1LHDfNPhtb/",
    "https://www.facebook.com/share/g/19NAzG4pYQ/",
    "https://www.facebook.com/share/g/1CdYu64paX/",
]

USER_FEED_MAX_SCROLLS = 400
USER_FEED_IDLE_LIMIT = 6
SCROLL_PAUSE_MS = 2_800
FALLBACK_MAX_SCROLLS = 600
FALLBACK_IDLE_LIMIT = 5

GAP_BETWEEN_GROUPS = (20, 45)
GAP_BETWEEN_POSTS = (4, 9)
LONG_PAUSE_EVERY = 15
LONG_PAUSE = (45, 90)
COMMENT_MAX_MORE = 60


def log(message: str) -> None:
    line = f"{datetime.now():%H:%M:%S} {message}"
    print(line, flush=True)
    try:
        LOG_FILE.parent.mkdir(parents=True, exist_ok=True)
        with LOG_FILE.open("a", encoding="utf-8") as handle:
            handle.write(line + "\n")
    except OSError:
        pass


# ------------------------------------------------------------- รายชื่อกลุ่ม
def load_groups(path: Path) -> list[str]:
    if not path.is_file():
        path.parent.mkdir(parents=True, exist_ok=True)
        path.write_text(
            "# กลุ่มที่จะไล่เก็บโพสต์ของตัวเอง — บรรทัดละ 1 ลิงก์\n"
            "# ขึ้นต้นด้วย # = ปิดบรรทัดนั้นไว้ชั่วคราว\n"
            + "\n".join(DEFAULT_GROUPS) + "\n",
            encoding="utf-8")
        log(f"สร้างรายชื่อกลุ่มไว้ที่ {path} ({len(DEFAULT_GROUPS)} กลุ่ม)")
    lines = [l.strip() for l in path.read_text(encoding="utf-8").splitlines()]
    return [l for l in lines if l and not l.startswith("#")]


def ensure_products_file(path: Path) -> None:
    """สร้างสมุดสินค้าเปล่าไว้ให้ — ผู้ใช้เติมเองทีหลังได้ ไม่เติมก็ยังทำงาน"""
    if path.is_file():
        return
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(
        "# สมุดสินค้า — ตั้งชื่อสินค้าเองได้ (ไม่กรอกก็ได้ ระบบจับกลุ่มจากข้อความให้เอง)\n"
        "#\n"
        "# รูปแบบ:   ชื่อสินค้า = คำค้น, คำค้น, คำค้น\n"
        "# เจอคำไหนในแคปชัน โพสต์นั้นจะถูกจัดเป็นสินค้านั้น\n"
        "#\n"
        "# ตัวอย่าง (ลบ # ข้างหน้าออกเพื่อใช้งาน):\n"
        "# ชาเขียวอเมซอน = ชาเขียว, amazon, ชาเย็น\n"
        "# หม้อทอดไร้น้ำมัน = หม้อทอด, air fryer, ไร้น้ำมัน\n"
        "# ที่พักทะเล = ที่พัก, รีสอร์ท, โรงแรม, ติดทะเล\n",
        encoding="utf-8")
    log(f"สร้างสมุดสินค้าเปล่าไว้ที่ {path} (เติมทีหลังได้)")


_GID_IN_URL = re.compile(r"/groups/(\d+)")


def resolve_group(page, url: str) -> dict:
    """คลี่ลิงก์ย่อเป็นเลขกลุ่ม + ชื่อกลุ่ม — ต้องอยู่ในเบราว์เซอร์ที่ล็อกอินแล้ว"""
    page.goto(url, wait_until="domcontentloaded", timeout=90_000)
    page.wait_for_timeout(4_000)
    if pc.guard_blocked(page):
        raise mf.MassFinderError("Facebook ขอยืนยันตัวตน (checkpoint) ตอนเปิดกลุ่ม")

    final = page.url or ""
    match = _GID_IN_URL.search(final)
    if not match:
        html = page.content()
        match = (re.search(r'"groupID"\s*:\s*"(\d+)"', html)
                 or re.search(r'"group_id"\s*:\s*"?(\d+)"?', html)
                 or _GID_IN_URL.search(html))
    if not match:
        raise mf.MassFinderError(f"คลี่ลิงก์กลุ่มไม่ได้: {url}")
    gid = match.group(1)

    name = ""
    for probe in ("h1", '[role="main"] h1', "title"):
        try:
            found = page.locator(probe).first
            if found.count():
                name = (found.inner_text(timeout=3_000) or "").strip()
                if name:
                    break
        except Exception:
            continue
    name = re.sub(r"\s*\|\s*Facebook\s*$", "", name).strip() or gid
    return {"gid": gid, "name": name, "share_url": url,
            "url": f"https://www.facebook.com/groups/{gid}/"}


# --------------------------------------------------------- เลื่อนเก็บโพสต์
def _scroll_collect(page, url: str, max_scrolls: int, idle_limit: int,
                    what: str) -> list[dict]:
    parser = pp.FeedParser()
    pending: list = []

    def on_response(response):
        if "/api/graphql" in response.url:
            pending.append(response)

    def drain() -> None:
        for response in pending:
            try:
                parser.add(response.text())
            except Exception:
                pass
        pending.clear()

    page.on("response", on_response)
    try:
        page.goto(url, wait_until="domcontentloaded", timeout=90_000)
        page.wait_for_timeout(5_000)
        pc.hush_page(page, log=log)
        if pc.guard_blocked(page):
            raise mf.MassFinderError("Facebook ขอยืนยันตัวตน (checkpoint)")
        drain()
        idle = 0
        for round_number in range(1, max_scrolls + 1):
            before = len(parser.posts)
            page.keyboard.press("End")
            page.wait_for_timeout(SCROLL_PAUSE_MS)
            drain()
            if round_number % 15 == 0:
                pc.hush_page(page, log=log)
                log(f"    {what}: เลื่อน {round_number} รอบ — เจอ {len(parser.posts)} โพสต์")
            idle = idle + 1 if len(parser.posts) == before else 0
            if idle >= idle_limit:
                log(f"    {what}: หมดแล้วที่ {len(parser.posts)} โพสต์ "
                    f"(เลื่อนไป {round_number} รอบ)")
                break
        else:
            log(f"    ⚠️ {what}: ชนเพดาน {max_scrolls} รอบ — **อาจยังไม่ครบ** "
                f"ได้มา {len(parser.posts)} โพสต์")
    finally:
        page.remove_listener("response", on_response)

    # `as_legacy()` ไม่ส่ง author_id ออกมา และลิงก์โปรไฟล์ที่เฟซส่งมามักเป็น
    # ชื่อเล่น (facebook.com/kp.oo.7) ซึ่งไม่มีเลขบัญชี — ถ้าไม่ดึงกลับมา
    # จะเหลือแต่การเทียบชื่อ ซึ่งพลาดทันทีที่มีเพจชื่อคล้ายกัน
    full = parser.result()
    rows = parser.as_legacy()
    for row in rows:
        row["author_id"] = str((full.get(row["id"]) or {}).get("author_id") or "")
    return rows


def mine_only(posts: list[dict], user_id: str, my_name: str) -> list[dict]:
    """คัดเฉพาะโพสต์ที่เราเป็นคนโพสต์

    เรียงความน่าเชื่อจากมากไปน้อย:
      1. `author_id` ตรงเป๊ะ                   ← แน่นอนที่สุด เปลี่ยนชื่อก็ยังเจอ
      2. เลขบัญชีโผล่ในลิงก์โปรไฟล์             ← ลิงก์แบบ profile.php?id=...
      3. ชื่อตรง **และ** ไม่มีตัวระบุอื่นเลย    ← ทางสุดท้าย สำหรับโพสต์เก่า
    """
    out = []
    for post in posts:
        author_id = str(post.get("author_id") or "")
        author_url = post.get("author_url") or ""
        author = (post.get("author") or "").strip()

        if author_id:
            if author_id == str(user_id):
                out.append(post)
            continue
        if user_id and (f"/{user_id}" in author_url or f"id={user_id}" in author_url):
            out.append(post)
            continue
        if not author_url and my_name and author.casefold() == my_name.casefold():
            out.append(post)
    return out


# ------------------------------------------------------------------ รูปภาพ
def save_images(post: dict, gid: str, log=log) -> list[str]:
    """โหลดรูปแนบทั้งหมดของโพสต์ — ลิงก์รูปเฟซมีวันหมดอายุ ต้องโหลดทันที"""
    folder = IMAGE_DIR / gid
    saved: list[str] = []
    for index, image in enumerate(post.get("images") or [], 1):
        uri = image.get("uri") if isinstance(image, dict) else str(image)
        if not uri:
            continue
        folder.mkdir(parents=True, exist_ok=True)
        target = folder / f"{post['id']}-{index}.jpg"
        if target.is_file() and target.stat().st_size > 0:
            saved.append(str(target))
            continue
        try:
            data = pc.fetch_image(uri)
            if not data:
                raise ValueError("ได้ไฟล์ว่าง")
            target.write_bytes(data)
            saved.append(str(target))
        except Exception as error:
            log(f"    ⚠️ โหลดรูปไม่ได้ {post['id']}-{index}: "
                f"{type(error).__name__} {error}")
    return saved


def save_comment_images(post_id: str, seq: int, urls: list, log=log) -> list[str]:
    """รูปในคอมเมนต์คือหลักฐานว่าลูกค้าถามอะไร — ทิ้งแล้วเหลือแต่ข้อความ"""
    folder = COMMENT_IMAGE_DIR / post_id
    saved: list[str] = []
    for index, uri in enumerate(urls or [], 1):
        if not uri:
            continue
        folder.mkdir(parents=True, exist_ok=True)
        target = folder / f"c{seq:04d}-{index}.jpg"
        if target.is_file() and target.stat().st_size > 0:
            saved.append(str(target))
            continue
        try:
            data = pc.fetch_image(uri)
            if not data:
                raise ValueError("ได้ไฟล์ว่าง")
            target.write_bytes(data)
            saved.append(str(target))
        except Exception as error:
            log(f"      ⚠️ รูปคอมเมนต์โหลดไม่ได้ {post_id} c{seq}-{index}: "
                f"{type(error).__name__}")
    return saved


# ------------------------------------------------------------ แต่งข้อมูลโพสต์
def finish_post(post: dict, group: dict, minimum: int) -> dict:
    """แปลงโพสต์ดิบเป็นแถวที่พร้อมวิเคราะห์ — ตัดสินแมส/ไม่แมส + คำนวณคุณสมบัติ"""
    likes = int(post.get("likes") or 0)
    comments = int(post.get("comments") or 0)
    shares = int(post.get("shares") or 0)
    total = likes + comments + shares
    unknown = post.get("unknown") or []

    # สามสถานะ ไม่ใช่สอง — โพสต์ที่ต่ำกว่าเส้นแต่มีช่องที่เฟซไม่ยอมบอกตัวเลข
    # **อาจเป็นแมสจริงก็ได้** ถ้าเหมาช่องที่ไม่รู้เป็นศูนย์แล้วตัดสินเลย
    # เท่ากับตัดสินผิดเงียบๆ ซึ่งจะทำให้บทวิเคราะห์ทั้งชุดเพี้ยน
    if total > minimum:
        status = "แมส"
    elif unknown:
        status = "ไม่แน่ใจ"
    else:
        status = "ไม่แมส"

    stamp = post.get("posted_at")
    row = {
        "group_name": group["name"], "gid": group["gid"],
        "id": post["id"], "url": post.get("url") or "",
        "caption": post.get("caption") or "",
        "likes": likes, "comments": comments, "shares": shares,
        "total": total, "unknown": unknown,
        "is_mass": status == "แมส", "mass_status": status,
        "image_urls": [i.get("uri") for i in (post.get("images") or [])
                       if isinstance(i, dict)],
        "image_files": [],
        "posted_at": stamp,
        "posted_at_text": (
            datetime.fromtimestamp(stamp).strftime("%Y-%m-%d %H:%M")
            if isinstance(stamp, (int, float)) and stamp else ""),
    }
    return row


def tag_products(rows: list[dict], book) -> dict:
    """ติดป้ายสินค้า/ชุดคอนเทนต์ให้ทุกโพสต์ — คืนข้อมูลชุดไว้ทำสรุป"""
    sets = feat.group_by_content(rows, book)
    for info in sets.values():
        for post in info["posts"]:
            post["product"] = info["label"]
            post["product_key"] = info["key"]
            post["product_from_book"] = info["from_book"]
            post["product_posts"] = info["n_posts"]
            post["product_groups"] = info["n_groups"]
    return sets


# ----------------------------------------------------------- เก็บคอมเมนต์
def collect_comments_for(page, rows: list[dict], minimum_comments: int,
                         on_progress=None, log=log) -> dict:
    """เปิดทีละโพสต์แล้วเก็บคอมเมนต์ให้ครบ — เขียนผลกลับเข้า row ทีละใบ

    **ข้ามโพสต์ที่ไม่มีคอมเมนต์เลย** ไม่ใช่การเก็บไม่ครบ — เปิดโพสต์ที่มี
    คอมเมนต์ 0 อันใช้เวลา 40 วินาทีเพื่อได้ของว่างเปล่า วัดจากกลุ่มจริง
    18 ส.ค.: 92 โพสต์ มีแค่ 5 ใบที่คอมเมนต์ > 0 = เสียเวลา 95% ไปเปล่าๆ
    """
    import random

    skipped = [r for r in rows
               if int(r.get("comments") or 0) < minimum_comments
               and not r.get("unknown")]
    for row in skipped:
        row["comments_done"] = True
        row["comment_list"] = []
        row["comment_skip"] = "ไม่มีคอมเมนต์ให้เก็บ"

    todo = [r for r in rows if not r.get("comments_done")]
    log(f"เก็บคอมเมนต์: ต้องเปิด {len(todo)} ใบ "
        f"· ข้าม {len(skipped)} ใบเพราะไม่มีคอมเมนต์ "
        f"(จากทั้งหมด {len(rows)} ใบ)")
    stats = {"posts": 0, "comments": 0, "images": 0, "incomplete": 0,
             "failed": 0, "skipped": len(skipped)}

    for number, row in enumerate(todo, 1):
        url = row.get("url") or ""
        if not url:
            row["comments_done"] = True
            row["comment_warn"] = "ไม่มีลิงก์โพสต์ เปิดไม่ได้"
            stats["failed"] += 1
            continue

        head = (row.get("caption") or "").strip().splitlines()
        head = (head[0][:36] if head else "(ไม่มีแคปชัน)")
        log(f"  [{number}/{len(todo)}] {row['mass_status']:8s} "
            f"{row['group_name'][:18]} · {head}")
        try:
            if pc.guard_blocked(page):
                raise mf.MassFinderError("Facebook ขอยืนยันตัวตน (checkpoint)")
            result = pcm.collect_comments(
                page, url, log=log, expect=int(row.get("comments") or 0),
                max_more=COMMENT_MAX_MORE)
        except mf.MassFinderError:
            raise
        except Exception as error:
            log(f"    ❌ เปิดไม่สำเร็จ: {type(error).__name__} {error}")
            row["comment_warn"] = f"{type(error).__name__}: {error}"
            stats["failed"] += 1
            if on_progress:
                on_progress()
            continue

        comments = []
        for seq, item in enumerate(result.get("comments") or [], 1):
            files = save_comment_images(row["id"], seq, item.get("images"),
                                        log=log)
            stats["images"] += len(files)
            text = item.get("text") or ""
            comments.append({
                "seq": seq, "author": item.get("author") or "",
                "author_url": item.get("author_url") or "",
                "avatar": item.get("avatar") or "",
                "when": item.get("when") or "", "text": text,
                "likes": int(item.get("likes") or 0),
                "depth": int(item.get("depth") or 0),
                "is_reply": bool(item.get("depth")),
                "image_urls": list(item.get("images") or []),
                "image_files": files,
                "is_spam": pp.is_spam(text), "spam_score": pp.spam_score(text),
            })

        row["comment_list"] = comments
        row["comments_done"] = True
        row["comments_got"] = result.get("got", len(comments))
        row["comments_expect"] = result.get("expect", 0)
        row["comment_warn"] = result.get("warn") or ""
        row["comment_seconds"] = result.get("seconds")
        stats["posts"] += 1
        stats["comments"] += len(comments)
        if result.get("warn"):
            stats["incomplete"] += 1

        if on_progress:
            on_progress()

        if number < len(todo):
            if number % LONG_PAUSE_EVERY == 0:
                nap = random.uniform(*LONG_PAUSE)
                log(f"    พักยาว {nap:.0f} วิ (ครบ {LONG_PAUSE_EVERY} ใบ)")
            else:
                nap = random.uniform(*GAP_BETWEEN_POSTS)
            time.sleep(nap)

    return stats


# ----------------------------------------------------------------- รายงาน
def _clip(text: str, limit: int = 900) -> str:
    text = (text or "").replace("\r", " ")
    return text if len(text) <= limit else text[:limit] + " …"


POST_COLUMNS = [
    ("กลุ่ม", 28), ("สินค้า/ชุดคอนเทนต์", 34), ("สถานะ", 10),
    ("วันที่โพสต์", 17), ("แคปชัน", 62),
    ("ไลค์", 8), ("คอมเมนต์", 10), ("แชร์", 8), ("รวม", 9),
    ("ตัวอักษร", 9), ("บรรทัด", 8), ("อิโมจิ", 8), ("แฮชแท็ก", 9),
    ("จำนวนรูป", 9), ("มีราคา", 8), ("การเปิดหัว", 18), ("ช่วงเวลา", 14),
    ("วัน", 8), ("แชร์ต่อ100ไลค์", 14), ("คอมเมนต์ที่เก็บได้", 17),
    ("ช่องที่อ่านไม่ได้", 16), ("ลิงก์โพสต์", 42), ("ไฟล์รูปในเครื่อง", 46),
]


def _post_row(row: dict) -> list:
    f = row.get("features") or {}
    got = row.get("comment_list")
    return [
        row.get("group_name"), row.get("product"), row.get("mass_status"),
        row.get("posted_at_text") or "", _clip(row.get("caption")),
        row.get("likes"), row.get("comments"), row.get("shares"),
        row.get("total"),
        f.get("ตัวอักษร"), f.get("บรรทัด"), f.get("อิโมจิ"), f.get("แฮชแท็ก"),
        f.get("จำนวนรูป"), "มี" if f.get("มีราคา") else "",
        f.get("การเปิดหัว"), f.get("ช่วงเวลา"), f.get("วัน"),
        f.get("แชร์ต่อ100ไลค์"),
        (len(got) if got is not None else "ยังไม่ได้เก็บ"),
        ", ".join(row.get("unknown") or []),
        row.get("url") or "", "\n".join(row.get("image_files") or []),
    ]


def write_excel(posts: list[dict], groups: list[dict], sets: dict,
                minimum: int, path: Path) -> None:
    """ออก Excel 6 ชีท — แมส · ไม่แมส · เทียบกัน · รายสินค้า · คอมเมนต์ · รายกลุ่ม"""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    book = Workbook()
    book.remove(book.active)
    head_font = Font(bold=True, color="FFFFFF")

    def sheet_of(title: str, columns, color="2F5496"):
        ws = book.create_sheet(title)
        ws.append([c for c, _ in columns])
        fill = PatternFill("solid", fgColor=color)
        for index, (_, width) in enumerate(columns, 1):
            ws.column_dimensions[get_column_letter(index)].width = width
            cell = ws.cell(row=1, column=index)
            cell.font, cell.fill = head_font, fill
            cell.alignment = Alignment(vertical="center")
        ws.freeze_panes = "A2"
        return ws

    def link_col(columns, title):
        return [c for c, _ in columns].index(title) + 1

    def fill(ws, items):
        cap = link_col(POST_COLUMNS, "แคปชัน")
        url = link_col(POST_COLUMNS, "ลิงก์โพสต์")
        files = link_col(POST_COLUMNS, "ไฟล์รูปในเครื่อง")
        for row in sorted(items, key=lambda r: r.get("total") or 0, reverse=True):
            ws.append(_post_row(row))
            last = ws.max_row
            for column in (cap, files):
                ws.cell(row=last, column=column).alignment = Alignment(
                    wrap_text=True, vertical="top")
            if row.get("url"):
                cell = ws.cell(row=last, column=url)
                cell.hyperlink, cell.style = row["url"], "Hyperlink"

    mass = [p for p in posts if p.get("mass_status") == "แมส"]
    plain = [p for p in posts if p.get("mass_status") == "ไม่แมส"]
    unsure = [p for p in posts if p.get("mass_status") == "ไม่แน่ใจ"]

    fill(sheet_of(f"แมส (เกิน {minimum})", POST_COLUMNS, "1F7A3D"), mass)
    fill(sheet_of("ไม่แมส", POST_COLUMNS, "8B3A2F"), plain)
    if unsure:
        fill(sheet_of("ไม่แน่ใจ (ตัวเลขไม่ครบ)", POST_COLUMNS, "8a6d00"), unsure)

    # ---- ชีทเทียบ: หัวใจของงานนี้ ----
    cmp_ws = sheet_of("เทียบ แมส vs ไม่แมส",
                      [("หมวด", 16), ("ช่อง", 26), ("แมส", 12),
                       ("ไม่แมส", 12), ("หน่วย", 10), ("ผลเทียบ", 24)],
                      "5B2C87")
    for line in feat.compare(mass, plain):
        cmp_ws.append([line["หมวด"], line["ช่อง"], line["แมส"],
                       line["ไม่แมส"], line["หน่วย"], line["ผลเทียบ"]])

    # ---- ชีทรายสินค้า ----
    prod = sheet_of("สรุปรายสินค้า",
                    [("สินค้า/ชุดคอนเทนต์", 42), ("ตั้งชื่อจากสมุด", 14),
                     ("ลงกี่โพสต์", 11), ("ลงกี่กลุ่ม", 11), ("แมสกี่ใบ", 10),
                     ("อัตราแมส %", 12), ("ยอดรวม", 11), ("ดีสุด", 10),
                     ("แย่สุด", 10), ("กลุ่มที่ลง", 60)], "0F6C74")
    for info in sorted(sets.values(), key=lambda s: s["total"], reverse=True):
        prod.append([
            info["label"], "ใช่" if info["from_book"] else "",
            info["n_posts"], info["n_groups"], info["n_mass"],
            round(info["n_mass"] * 100 / info["n_posts"], 1) if info["n_posts"] else 0,
            info["total"], info["best"], info["worst"],
            ", ".join(info["groups"]),
        ])

    # ---- ชีทคอมเมนต์ ----
    all_comments = [(p, c) for p in posts for c in (p.get("comment_list") or [])]
    if all_comments:
        cs = sheet_of("คอมเมนต์ทั้งหมด", [
            ("กลุ่ม", 24), ("สินค้า/ชุดคอนเทนต์", 30), ("สถานะโพสต์", 10),
            ("ลำดับ", 7), ("เป็นคำตอบ", 10), ("คนพิมพ์", 20),
            ("ข้อความ", 64), ("ถูกใจ", 8), ("เมื่อไร", 13), ("รูปแนบ", 8),
            ("สแปม", 8), ("ลิงก์โพสต์", 42)], "444444")
        text_col = 7
        for post, item in all_comments:
            cs.append([
                post["group_name"], post.get("product"), post.get("mass_status"),
                item.get("seq"), "ใช่" if item.get("is_reply") else "",
                item.get("author"), _clip(item.get("text"), 700),
                item.get("likes"), item.get("when"),
                len(item.get("image_files") or []),
                "สแปม" if item.get("is_spam") else "",
                post.get("url") or "",
            ])
            cs.cell(row=cs.max_row, column=text_col).alignment = Alignment(
                wrap_text=True, vertical="top")

    # ---- ชีทรายกลุ่ม ----
    summary = sheet_of("สรุปรายกลุ่ม",
                       [("กลุ่ม", 38), ("เลขกลุ่ม", 18), ("โพสต์ของเรา", 14),
                        ("แมส", 8), ("ไม่แมส", 10), ("ไม่แน่ใจ", 10),
                        ("อัตราแมส %", 12), ("ยอดรวม", 12),
                        ("วิธีที่ใช้เก็บ", 24)], "2F5496")
    for group in groups:
        summary.append([
            group.get("name"), group.get("gid"), group.get("mine_found", 0),
            group.get("mass", 0), group.get("plain", 0), group.get("unsure", 0),
            group.get("mass_rate", 0), group.get("total_engagement", 0),
            group.get("how", ""),
        ])

    path.parent.mkdir(parents=True, exist_ok=True)
    book.save(path)


# -------------------------------------------------------------------- หลัก
def collect(bot_name: str, group_urls: list[str], minimum: int,
            with_comments: bool = True,
            comments_min: int = COMMENTS_MIN) -> dict:
    from playwright.sync_api import sync_playwright

    farm = ProfileFarm(mf.DATA_DIR)
    entry = mf.find_bot(farm, bot_name)

    ensure_products_file(PRODUCTS_FILE)
    book = feat.load_products(PRODUCTS_FILE)
    if book:
        log(f"สมุดสินค้า: {len(book)} รายการ ({', '.join(n for n, _ in book[:5])}…)")

    OUT_DIR.mkdir(parents=True, exist_ok=True)
    posts: list[dict] = []
    group_stats: list[dict] = []

    with sync_playwright() as pw:
        context = mf.launch_bot_browser(pw, farm, entry)
        try:
            context.add_init_script(pc.MUTE_JS)
            page = context.pages[0] if context.pages else context.new_page()
            pc.hush_page(page, log=log)

            user_id = mf.ensure_logged_in(page, context)
            my_name = ""
            try:
                page.goto(f"https://www.facebook.com/{user_id}",
                          wait_until="domcontentloaded", timeout=60_000)
                page.wait_for_timeout(3_000)
                my_name = (page.locator("h1").first.inner_text(timeout=5_000)
                           or "").strip()
            except Exception:
                pass
            log(f"ล็อกอินเป็น: {my_name or '(อ่านชื่อไม่ได้)'} · เลขบัญชี {user_id}")

            for number, share_url in enumerate(group_urls, 1):
                log(f"[{number}/{len(group_urls)}] กลุ่ม {share_url}")
                try:
                    group = resolve_group(page, share_url)
                except Exception as error:
                    log(f"  ❌ ข้ามกลุ่มนี้: {type(error).__name__} {error}")
                    group_stats.append({"name": share_url, "gid": "-",
                                        "how": f"ล้มเหลว: {error}"})
                    continue
                log(f"  = {group['name']} (เลขกลุ่ม {group['gid']})")

                how = "หน้าโพสต์ของเราโดยตรง"
                user_feed = (f"https://www.facebook.com/groups/"
                             f"{group['gid']}/user/{user_id}/")
                raw = _scroll_collect(page, user_feed, USER_FEED_MAX_SCROLLS,
                                      USER_FEED_IDLE_LIMIT, "โพสต์ของเรา")
                mine = mine_only(raw, user_id, my_name)

                if not mine:
                    log("  ⚠️ หน้าโพสต์ของเราไม่ให้ข้อมูล — ถอยไปไล่ทั้งกลุ่มแทน "
                        "(ช้ากว่ามาก และอาจไม่ครบถ้าโพสต์เก่ากว่าเพดานเลื่อน)")
                    how = "ไล่ทั้งกลุ่มแล้วกรอง (ทางถอย)"
                    raw = _scroll_collect(page, group["url"],
                                          FALLBACK_MAX_SCROLLS,
                                          FALLBACK_IDLE_LIMIT, "ทั้งกลุ่ม")
                    mine = mine_only(raw, user_id, my_name)

                rows = [finish_post(p, group, minimum) for p in mine]
                by_id = {p["id"]: p for p in mine}
                for row in rows:
                    row["image_files"] = save_images(by_id[row["id"]],
                                                     group["gid"], log=log)
                    row["features"] = feat.extract(row)
                posts.extend(rows)

                n_mass = sum(1 for r in rows if r["mass_status"] == "แมส")
                n_plain = sum(1 for r in rows if r["mass_status"] == "ไม่แมส")
                n_unsure = sum(1 for r in rows if r["mass_status"] == "ไม่แน่ใจ")
                group_stats.append({
                    "name": group["name"], "gid": group["gid"],
                    "mine_found": len(rows), "mass": n_mass, "plain": n_plain,
                    "unsure": n_unsure,
                    "mass_rate": round(n_mass * 100 / len(rows), 1) if rows else 0,
                    "total_engagement": sum(r["total"] for r in rows),
                    "how": how,
                })
                log(f"  โพสต์ของเรา {len(rows)} ใบ → แมส {n_mass} · "
                    f"ไม่แมส {n_plain} · ไม่แน่ใจ {n_unsure}")

                sets = tag_products(posts, book)
                _save_raw(posts, group_stats, sets, minimum, user_id, my_name)

                if number < len(group_urls):
                    import random
                    nap = random.uniform(*GAP_BETWEEN_GROUPS)
                    log(f"  พัก {nap:.0f} วิ ก่อนกลุ่มถัดไป")
                    time.sleep(nap)

            sets = tag_products(posts, book)
            log("─" * 60)
            log(f"รวมทุกกลุ่ม: {len(posts)} โพสต์ · "
                f"{len(sets)} ชุดคอนเทนต์/สินค้า")

            if with_comments and posts:
                stats = collect_comments_for(
                    page, posts, comments_min,
                    on_progress=lambda: _save_raw(posts, group_stats, sets,
                                                  minimum, user_id, my_name),
                    log=log)
                log(f"คอมเมนต์: {stats['comments']:,} อัน จาก {stats['posts']} ใบ "
                    f"· รูปในคอมเมนต์ {stats['images']} รูป")
                if stats["incomplete"]:
                    log(f"⚠️ {stats['incomplete']} ใบเก็บได้ไม่ครบ")
                if stats["failed"]:
                    log(f"⚠️ {stats['failed']} ใบเปิดไม่สำเร็จเลย")
                _save_raw(posts, group_stats, sets, minimum, user_id, my_name)
        finally:
            try:
                context.close()
            except Exception:
                pass

    return {"posts": posts, "groups": group_stats, "sets": sets}


def _save_raw(posts, groups, sets, minimum, user_id, my_name) -> None:
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    slim = {k: {kk: vv for kk, vv in v.items() if kk != "posts"}
            for k, v in (sets or {}).items()}
    RAW_FILE.write_text(json.dumps({
        "saved_at": datetime.now().isoformat(timespec="seconds"),
        "user_id": user_id, "my_name": my_name, "minimum": minimum,
        "groups": groups, "sets": slim, "posts": posts,
    }, ensure_ascii=False, indent=1), encoding="utf-8")


def preflight(bot_name: str, group_urls: list[str]) -> int:
    ok = True
    log(f"กลุ่มที่จะเก็บ: {len(group_urls)} กลุ่ม")
    for url in group_urls:
        log(f"   · {url}")
    if not group_urls:
        log("❌ ไม่มีกลุ่มให้เก็บเลย")
        ok = False

    try:
        farm = ProfileFarm(mf.DATA_DIR)
        entry = mf.find_bot(farm, bot_name)
        log(f"บอท: {entry['name']} ({entry['id']}) · Chrome เปิดค้างอยู่ไหม: "
            f"{'ใช่ ← ต้องปิดก่อน' if entry.get('running') else 'ไม่'}")
    except Exception as error:
        log(f"❌ บอท: {error}")
        ok = False

    ensure_products_file(PRODUCTS_FILE)
    book = feat.load_products(PRODUCTS_FILE)
    log(f"สมุดสินค้า: {len(book)} รายการ"
        + ("" if book else " (ว่าง — ระบบจะจับกลุ่มจากตัวข้อความให้เอง)"))

    for name, module in (("openpyxl", "openpyxl"), ("playwright", "playwright")):
        try:
            __import__(module)
            log(f"{name}: มี")
        except ImportError:
            log(f"❌ {name}: ไม่มี")
            ok = False

    log(f"ที่เก็บผล: {OUT_DIR}")
    log("✅ ตรวจผ่าน พร้อมรัน" if ok else "❌ ยังไม่พร้อม แก้ตามข้างบนก่อน")
    return 0 if ok else 1


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--bot", default=DEFAULT_BOT)
    parser.add_argument("--mass", type=int, default=MASS_MIN,
                        help="ไลค์+คอมเมนต์+แชร์ เกินเท่าไรถือว่าแมส (ค่าตั้งต้น 100)")
    parser.add_argument("--comments-min", type=int, default=COMMENTS_MIN,
                        help="เปิดเก็บคอมเมนต์เมื่อโพสต์มีคอมเมนต์ตั้งแต่กี่อัน")
    parser.add_argument("--groups-file", default=str(GROUPS_FILE))
    parser.add_argument("--dry-run", action="store_true")
    parser.add_argument("--no-comments", action="store_true")
    args = parser.parse_args(argv)

    group_urls = load_groups(Path(args.groups_file))
    if args.dry_run:
        return preflight(args.bot, group_urls)
    if preflight(args.bot, group_urls) != 0:
        return 1

    log("─" * 60)
    started = time.time()
    try:
        result = collect(args.bot, group_urls, args.mass,
                         with_comments=not args.no_comments,
                         comments_min=args.comments_min)
    except mf.MassFinderError as error:
        log(f"❌ {error}")
        return 2
    except Exception as error:
        log(f"❌ พังกลางทาง: {type(error).__name__} {error}")
        log(traceback.format_exc())
        return 3

    posts, sets = result["posts"], result["sets"]
    write_excel(posts, result["groups"], sets, args.mass, XLSX_FILE)

    mass = sum(1 for p in posts if p["mass_status"] == "แมส")
    plain = sum(1 for p in posts if p["mass_status"] == "ไม่แมส")
    unsure = sum(1 for p in posts if p["mass_status"] == "ไม่แน่ใจ")
    comments = sum(len(p.get("comment_list") or []) for p in posts)
    images = sum(len(p.get("image_files") or []) for p in posts)

    log("─" * 60)
    log(f"เสร็จใน {(time.time()-started)/60:.1f} นาที")
    log(f"โพสต์ทั้งหมด {len(posts)} ใบ → แมส {mass} · ไม่แมส {plain}"
        + (f" · ไม่แน่ใจ {unsure}" if unsure else ""))
    log(f"ชุดคอนเทนต์/สินค้า: {len(sets)} ชุด")
    log(f"คอมเมนต์ {comments:,} อัน · รูปในโพสต์ {images} รูป")
    log(f"Excel     : {XLSX_FILE}")
    log(f"ข้อมูลดิบ : {RAW_FILE}")
    log("ขั้นต่อไป : python fb_myposts_view.py  (สร้างหน้าเว็บแบบเฟซบุ๊ก)")
    return 0


if __name__ == "__main__":
    sys.exit(main())
