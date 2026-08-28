"""ทดสอบตัวเก็บโพสต์ตัวเอง — ส่วนที่ทดสอบได้โดยไม่ต้องเปิดเบราว์เซอร์

ครอบเรื่องที่พลาดแล้วเสียหายจริง:
  1. คัดโพสต์ผิดคน          → เอาข้อมูลคนอื่นมาวิเคราะห์
  2. ตัดสินแมส/ไม่แมสเพี้ยน → บทวิเคราะห์ทั้งชุดผิดตาม
  3. ตัวเลขอ่านไม่ได้ถูกเหมาเป็น 0 → โพสต์ดังหล่นไปกองไม่แมสเงียบๆ
  4. จับกลุ่มสินค้าผิด      → เทียบข้ามสินค้ากันโดยไม่รู้ตัว

รัน:  python test_myposts.py
"""
from __future__ import annotations

import os
import sys
import tempfile
from pathlib import Path

os.environ.setdefault("STUDIO_DATA_DIR",
                      str(Path(tempfile.gettempdir()) / "myposts-test-data"))

import fb_myposts as m              # noqa: E402
import fb_myposts_features as feat  # noqa: E402

PASS = FAIL = 0


def check(name: str, got, want) -> None:
    global PASS, FAIL
    if got == want:
        PASS += 1
        print(f"  ✅ {name}")
    else:
        FAIL += 1
        print(f"  ❌ {name}\n       ได้  : {got!r}\n       ควรได้: {want!r}")


ME = "61550000000001"
MY_NAME = "Kp Oo"
GROUP = {"gid": "111", "name": "กลุ่มทดสอบ"}


def post(pid, author="", author_url="", author_id="", likes=0, comments=0,
         shares=0, unknown=(), caption="", images=(), when=None):
    return {"id": pid, "url": f"https://www.facebook.com/groups/1/posts/{pid}/",
            "author": author, "author_url": author_url, "author_id": author_id,
            "likes": likes, "comments": comments, "shares": shares,
            "unknown": list(unknown), "caption": caption,
            "images": [{"uri": u, "alt": ""} for u in images],
            "posted_at": when}


# ------------------------------------------------------------------ 1-3
print("\n1) คัดเฉพาะโพสต์ของเรา")
feed = [
    post("p1", MY_NAME, f"https://www.facebook.com/{ME}"),
    post("p2", "คนอื่น", "https://www.facebook.com/61559999999999"),
    post("p3", MY_NAME, f"https://www.facebook.com/profile.php?id={ME}"),
    post("p4", MY_NAME, ""),
    post("p5", "Kp Oo Fanpage", "https://www.facebook.com/61551111111111"),
]
mine = m.mine_only(feed, ME, MY_NAME)
check("ได้เฉพาะของเรา", sorted(p["id"] for p in mine), ["p1", "p3", "p4"])
check("ชื่อคล้ายกันแต่คนละบัญชี ไม่ติดมา",
      any(p["id"] == "p5" for p in mine), False)

print("\n2) ลิงก์เป็นชื่อเล่น — ของจริงเฟซส่งมาแบบนี้")
nick = [
    post("n1", MY_NAME, "https://www.facebook.com/kp.oo.7", author_id=ME),
    post("n3", MY_NAME, "https://www.facebook.com/impostor",
         author_id="61558888888888"),
]
check("จับจากเลขบัญชีได้", [p["id"] for p in m.mine_only(nick, ME, MY_NAME)],
      ["n1"])
check("ชื่อตรงแต่เลขคนละคน = ไม่ใช่ของเรา",
      any(p["id"] == "n3" for p in m.mine_only(nick, ME, MY_NAME)), False)

print("\n3) ไม่มีอะไรให้ยึดเลย = ไม่เดา")
check("ไม่เดาว่าเป็นของเรา", m.mine_only([post("x", "สมชาย", "")], ME, MY_NAME), [])

# -------------------------------------------------------------------- 4
print("\n4) ตัดสินแมส / ไม่แมส / ไม่แน่ใจ")
cases = [
    (dict(likes=101), "แมส", "ไลค์ 101 = แมส"),
    (dict(likes=100), "ไม่แมส", "ไลค์ 100 เป๊ะ = ยังไม่แมส"),
    (dict(likes=50, comments=30, shares=21), "แมส", "50+30+21 = 101 แมส"),
    (dict(likes=50, comments=30, shares=20), "ไม่แมส", "50+30+20 = 100 ไม่แมส"),
    (dict(likes=90, unknown=["shares"]), "ไม่แน่ใจ",
     "ต่ำกว่าเส้นแต่แชร์อ่านไม่ได้ = ไม่แน่ใจ"),
    (dict(likes=200, unknown=["shares"]), "แมส",
     "เกินเส้นอยู่แล้ว แม้แชร์อ่านไม่ได้ = แมสแน่นอน"),
    (dict(), "ไม่แมส", "ยอด 0 ทุกช่อง อ่านได้ครบ = ไม่แมส"),
]
for kwargs, want, label in cases:
    row = m.finish_post(post("t", **kwargs), GROUP, 100)
    check(label, row["mass_status"], want)

row = m.finish_post(post("t", likes=150), GROUP, 100)
check("is_mass ตรงกับ mass_status", row["is_mass"], True)
check("ยอดรวมคิดถูก",
      m.finish_post(post("t", likes=10, comments=5, shares=2), GROUP, 100)["total"], 17)

# -------------------------------------------------------------------- 5
print("\n5) แกะคุณสมบัติจากแคปชัน")
cap = ("ราคานี้จริงเหรอ 😱 ที่พักติดทะเล คืนละ 690 บาท\n"
       "รวมอาหารเช้า เหลือไม่กี่ห้อง ทักแชทได้เลย\n"
       "#ที่พัก #รีวิว https://example.com/x")
f = feat.extract(m.finish_post(post("z", caption=cap, likes=100, shares=25,
                                    comments=50), GROUP, 100))
check("นับบรรทัดถูก", f["บรรทัด"], 3)
check("นับแฮชแท็กถูก", f["แฮชแท็ก"], 2)
check("นับลิงก์ถูก", f["ลิงก์"], 1)
check("จับได้ว่ามีราคา", f["มีราคา"], 1)
check("จับอิโมจิได้", f["อิโมจิ"] >= 1, True)
check("จับคำเร่งด่วนได้", f["เร่งด่วน"], 1)
check("จับคำชวนทักได้", f["ชวนทัก"], 1)
check("แชร์ต่อ100ไลค์", f["แชร์ต่อ100ไลค์"], 25.0)
check("คอมเมนต์ต่อ100ไลค์", f["คอมเมนต์ต่อ100ไลค์"], 50.0)

print("\n6) แยกประเภทการเปิดหัว")
for text, want in [
    ("ราคา 390 บาทเท่านั้น\nรายละเอียด", "เปิดด้วยราคา"),
    ("ใครเคยลองบ้างคะ\nอยากรู้", "เปิดด้วยคำถาม"),
    ("โอ้โห ของมันต้องมี!\nรีบเลย", "เปิดด้วยอุทาน"),
    ("สวัสดีค่ะ ขออนุญาตฝากร้าน", "ทักทาย/ขออนุญาต"),
    ("วันนี้ไปเที่ยวทะเลมา บรรยากาศดีมาก", "เปิดด้วยเล่าเรื่อง"),
    ("", "ไม่มีข้อความ"),
]:
    check(f"{(text[:22] or '(ว่าง)')}…", feat.opener_kind(text), want)

print("\n7) แบ่งช่วงเวลา")
for hour, want in [(2, "ดึก (00-06)"), (8, "เช้า (06-11)"), (12, "เที่ยง (11-14)"),
                   (15, "บ่าย (14-17)"), (19, "เย็น (17-21)"), (22, "ค่ำ (21-24)")]:
    check(f"{hour} นาฬิกา", feat.time_slot(hour), want)

# -------------------------------------------------------------------- 8
print("\n8) จับกลุ่มสินค้า — โพสต์เดียวกันลงหลายกลุ่มต้องอยู่ชุดเดียว")
same = "ชาเขียวอเมซอน แก้วละ {p} บาท 😳 อร่อยมาก #รีวิว"
rows = []
for i, (gid, gname, price) in enumerate([("1", "กลุ่ม ก", 39),
                                         ("2", "กลุ่ม ข", 45),
                                         ("3", "กลุ่ม ค", 39)]):
    r = m.finish_post(post(f"s{i}", caption=same.format(p=price), likes=150),
                      {"gid": gid, "name": gname}, 100)
    r["features"] = feat.extract(r)
    rows.append(r)
other = m.finish_post(post("o1", caption="หม้อทอดไร้น้ำมัน ลดเหลือ 990 บาท",
                           likes=20), GROUP, 100)
other["features"] = feat.extract(other)
rows.append(other)

sets = m.tag_products(rows, [])
check("ได้ 2 ชุด (ชาเขียว 1 · หม้อทอด 1)", len(sets), 2)
tea = [s for s in sets.values() if "ชาเขียว" in s["label"]][0]
check("ชาเขียวรวมได้ 3 โพสต์ แม้ราคาต่างกัน", tea["n_posts"], 3)
check("ชาเขียวลงไป 3 กลุ่ม", tea["n_groups"], 3)
check("ทุกโพสต์ได้ป้ายสินค้า", all(r.get("product") for r in rows), True)
check("ชาเขียวสามใบได้ป้ายเดียวกัน",
      len({r["product"] for r in rows[:3]}), 1)

print("\n9) สมุดสินค้าที่ผู้ใช้ตั้งชื่อเอง ชนะชื่ออัตโนมัติ")
bookfile = Path(tempfile.mkdtemp()) / "products.txt"
bookfile.write_text("# หมายเหตุ\nชาเขียวอเมซอน = ชาเขียว, amazon\n"
                    "หม้อทอด = หม้อทอด, air fryer\nบรรทัดผิด\n",
                    encoding="utf-8")
bk = feat.load_products(bookfile)
check("อ่านสมุดได้ 2 รายการ ข้ามบรรทัดผิด", len(bk), 2)
sets2 = m.tag_products(rows, bk)
labels = sorted({s["label"] for s in sets2.values()})
check("ใช้ชื่อจากสมุด", labels, ["ชาเขียวอเมซอน", "หม้อทอด"])
check("ทำเครื่องหมายว่ามาจากสมุด",
      all(s["from_book"] for s in sets2.values()), True)
check("ไม่ตรงสมุดเลย = ไม่เดาชื่อ", feat.match_product("ข้อความอื่น", bk), "")

# ------------------------------------------------------------------- 10
print("\n10) ตารางเทียบ แมส vs ไม่แมส")
mass = [dict(features=feat.extract(m.finish_post(
    post(f"m{i}", caption="ราคา 390 บาท ทักแชทเลย 😱", likes=200), GROUP, 100)))
    for i in range(3)]
plain = [dict(features=feat.extract(m.finish_post(
    post(f"n{i}", caption="วันนี้อากาศดี", likes=5), GROUP, 100)))
    for i in range(3)]
table = feat.compare(mass, plain)
check("ได้ตารางเทียบ", len(table) > 10, True)
row_price = [r for r in table if r["ช่อง"] == "มีราคา"][0]
check("แมสมีราคา 100%", row_price["แมส"], 100.0)
check("ไม่แมสไม่มีราคา 0%", row_price["ไม่แมส"], 0.0)
check("บอกผลเทียบเป็นภาษาคน", row_price["ผลเทียบ"], "แมสมี ไม่แมสไม่มี")
row_len = [r for r in table if r["ช่อง"] == "ตัวอักษร"][0]
check("แมสยาวกว่า", row_len["แมส"] > row_len["ไม่แมส"], True)
check("เทียบกับกลุ่มว่างได้ ไม่พัง", len(feat.compare(mass, [])) > 0, True)

# ------------------------------------------------------------------- 11
print("\n11) ออก Excel ครบทุกชีท")
out = Path(tempfile.mkdtemp()) / "t.xlsx"
demo = rows[:]
for i, r in enumerate(demo):
    r["comment_list"] = [{
        "seq": 1, "author": "สมหญิง", "text": "ราคาเท่าไรคะ", "likes": 12,
        "when": "2 สัปดาห์", "is_reply": False, "image_files": [], "is_spam": False,
    }, {
        "seq": 2, "author": "#ปล่อยกู้", "text": "#ดอก29", "likes": 0,
        "when": "1 วัน", "is_reply": True, "image_files": ["C:/x.jpg"],
        "is_spam": True,
    }]
gstats = [{"name": "กลุ่ม ก", "gid": "1", "mine_found": 3, "mass": 3,
           "plain": 0, "unsure": 0, "mass_rate": 100.0,
           "total_engagement": 450, "how": "หน้าโพสต์ของเราโดยตรง"}]
m.write_excel(demo, gstats, sets, 100, out)

from openpyxl import load_workbook  # noqa: E402
book = load_workbook(out)
check("มีชีทแมส", "แมส (เกิน 100)" in book.sheetnames, True)
check("มีชีทไม่แมส", "ไม่แมส" in book.sheetnames, True)
check("มีชีทเทียบ", "เทียบ แมส vs ไม่แมส" in book.sheetnames, True)
check("มีชีทรายสินค้า", "สรุปรายสินค้า" in book.sheetnames, True)
check("มีชีทคอมเมนต์", "คอมเมนต์ทั้งหมด" in book.sheetnames, True)
check("มีชีทรายกลุ่ม", "สรุปรายกลุ่ม" in book.sheetnames, True)


def col_of(sheet, title: str) -> int:
    """หาช่องจากชื่อหัวตาราง ไม่ใช่ตำแหน่ง — เพิ่มช่องใหม่แล้วเทสไม่พังตาม"""
    for cell in sheet[1]:
        if cell.value == title:
            return cell.column
    raise AssertionError(f"ไม่พบหัวตาราง {title!r}")


ws = book["แมส (เกิน 100)"]
check("แถวโพสต์แมสครบ", ws.max_row - 1, 3)
check("มีช่องสินค้า", ws.cell(row=2, column=col_of(ws, "สินค้า/ชุดคอนเทนต์")).value,
      "ชาเขียวอเมซอน")
check("มีช่องสถานะ", ws.cell(row=2, column=col_of(ws, "สถานะ")).value, "แมส")
check("ลิงก์คลิกได้",
      ws.cell(row=2, column=col_of(ws, "ลิงก์โพสต์")).hyperlink is not None, True)
check("มีช่องคุณสมบัติ (การเปิดหัว)",
      ws.cell(row=2, column=col_of(ws, "การเปิดหัว")).value is not None, True)

cs = book["คอมเมนต์ทั้งหมด"]
check("คอมเมนต์ครบทุกอัน (4 โพสต์ × 2 อัน)", cs.max_row - 1, len(demo) * 2)
check("คอมเมนต์รู้ว่ามาจากสินค้าไหน",
      cs.cell(row=2, column=col_of(cs, "สินค้า/ชุดคอนเทนต์")).value,
      "ชาเขียวอเมซอน")
check("เก็บยอดถูกใจรายคอมเมนต์",
      cs.cell(row=2, column=col_of(cs, "ถูกใจ")).value, 12)
check("ทำเครื่องหมายสแปม",
      cs.cell(row=3, column=col_of(cs, "สแปม")).value, "สแปม")

ps = book["สรุปรายสินค้า"]
check("สรุปรายสินค้าครบ", ps.max_row - 1, 2)
check("บอกว่าลงกี่กลุ่ม",
      ps.cell(row=2, column=col_of(ps, "ลงกี่กลุ่ม")).value, 3)

print("\n12) ไม่มีโพสต์ 'ไม่แน่ใจ' = ไม่ต้องมีชีทนั้น")
out2 = Path(tempfile.mkdtemp()) / "t2.xlsx"
m.write_excel([r for r in demo if r["mass_status"] != "ไม่แน่ใจ"],
              gstats, sets, 100, out2)
check("ไม่มีชีทไม่แน่ใจ",
      "ไม่แน่ใจ (ตัวเลขไม่ครบ)" in load_workbook(out2).sheetnames, False)

print("\n13) อ่านรายชื่อกลุ่มจากไฟล์")
tmp = Path(tempfile.mkdtemp()) / "groups.txt"
tmp.write_text("# หมายเหตุ\nhttps://a/\n\n  https://b/  \n#https://c/\n",
               encoding="utf-8")
check("ตัดคอมเมนต์กับบรรทัดว่างออก", m.load_groups(tmp), ["https://a/", "https://b/"])
check("ไม่มีไฟล์ = สร้างให้พร้อม 6 กลุ่ม",
      len(m.load_groups(Path(tempfile.mkdtemp()) / "new.txt")), 6)

print(f"\n{'=' * 46}\nผ่าน {PASS} · ล้ม {FAIL}")
sys.exit(1 if FAIL else 0)
